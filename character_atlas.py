"""
Character Atlas - ComfyUI node
Builds a character description from character_atlas_library.xlsx and prepends it to a prompt.
"""
import os
import glob
import random

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

HERE = os.path.dirname(os.path.abspath(__file__))
NONE = "None"
RAND = "🎲 Random"
SEP = " | "


# --------------------------------------------------------------------------- library
def _find_library():
    p = os.path.join(HERE, "character_atlas_library.xlsx")
    if os.path.exists(p):
        return p
    found = sorted(glob.glob(os.path.join(HERE, "character_atlas_library*.xlsx")))
    return found[0] if found else None


def _read_sheet(wb, name):
    if name not in wb.sheetnames:
        return []
    rows = wb[name].iter_rows(values_only=True)
    try:
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return []
    out = []
    for r in rows:
        if r is None or all(v is None for v in r):
            continue
        d = {header[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(header)}
        if d.get("Enabled", "Yes").lower() in ("no", "false", "0"):
            continue
        out.append(d)
    return out


class Library:
    def __init__(self):
        self.path = _find_library()
        self.error = None
        self.known, self.presets, self.archetypes, self.attrs, self.outfits = {}, {}, {}, {}, {}
        self.colors, self.fabrics, self.patterns, self.finishes = {}, {}, {}, {}
        self.poses, self.accessories, self.expressions = {}, {}, {}
        if openpyxl is None:
            self.error = "openpyxl not installed"
            return
        if not self.path:
            self.error = "character_atlas_library.xlsx not found next to character_atlas.py"
            return
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        for r in _read_sheet(wb, "Known_Characters"):
            self.known[f"{r['Franchise']}{SEP}{r['Character']}"] = r
        for r in _read_sheet(wb, "Presets"):
            self.presets[f"{r['Franchise']}{SEP}{r['Role']} ({r['Actor']})"] = r
        for r in _read_sheet(wb, "Archetypes"):
            self.archetypes[f"{r['Category']}{SEP}{r['Name']}"] = r
        for r in _read_sheet(wb, "Attributes"):
            t = r["Attribute_Type"]
            key = f"{r['Category']}{SEP}{r['Name']}" if t in ("Body Shape", "Hair Style") else r["Name"]
            self.attrs.setdefault(t, {})[key] = r
        for r in _read_sheet(wb, "Outfits"):
            self.outfits.setdefault(r["Slot"], {})[f"{r['Category']}{SEP}{r['Name']}"] = r
        for r in _read_sheet(wb, "Colors"):
            self.colors[f"{r['Category']}{SEP}{r['Name']}"] = r
        for r in _read_sheet(wb, "Fabrics"):
            self.fabrics[f"{r['Category']}{SEP}{r['Name']}"] = r
        for r in _read_sheet(wb, "Patterns"):
            self.patterns[r["Name"]] = r
        for r in _read_sheet(wb, "Finishes"):
            self.finishes[r["Name"]] = r
        for r in _read_sheet(wb, "Expressions"):
            self.expressions[f"{r['Category']}{SEP}{r['Name']}"] = r
        for r in _read_sheet(wb, "Poses"):
            self.poses[f"{r['Category']}{SEP}{r['Subcategory']}{SEP}{r['Pose']}"] = r
        for r in _read_sheet(wb, "Accessories"):
            self.accessories[f"{r['Category']}{SEP}{r['Name']}"] = r
        wb.close()

    def fabrics_for(self, slot):
        return {k: r for k, r in self.fabrics.items()
                if slot in [s.strip() for s in r.get("Applies_To", "").split(";")]}


LIB = Library()
if LIB.error:
    print(f"[Character Atlas] WARNING: {LIB.error}")
else:
    print(f"[Character Atlas] loaded {len(LIB.known)} known characters, {len(LIB.presets)} presets, "
          f"{len(LIB.archetypes)} archetypes, {len(LIB.expressions)} expressions, {len(LIB.poses)} poses, "
          f"{len(LIB.accessories)} accessories from {os.path.basename(LIB.path)}")


# --------------------------------------------------------------------------- random-aware options
SOURCES = {}  # widget name -> {key: row}


def _opts(name, mapping, randomizable=True):
    """Dropdown list: None, 🎲 Random, then entries with '<Category> | 🎲 Random' before each group."""
    SOURCES[name] = mapping
    keys = list(mapping.keys())
    if not randomizable:
        return [NONE] + keys
    out, seen = [NONE, RAND], set()
    for k in keys:
        parts = k.split(SEP)
        for depth in range(1, len(parts)):
            prefix = SEP.join(parts[:depth])
            if prefix in seen:
                continue
            seen.add(prefix)
            if sum(1 for x in keys if x.startswith(prefix + SEP)) > 1:
                out.append(f"{prefix}{SEP}{RAND}")
        out.append(k)
    return out


def _resolve(name, value, seed):
    """Turn a random token into a concrete key (deterministic per seed + widget)."""
    if value in (None, NONE) or RAND not in str(value):
        return value
    keys = list(SOURCES.get(name, {}).keys())
    if value != RAND:
        prefix = value[: -len(SEP + RAND)]
        keys = [k for k in keys if k.startswith(prefix + SEP)]
    if not keys:
        return NONE
    return random.Random(f"{seed}:{name}").choice(keys)


# --------------------------------------------------------------------------- text helpers
def _pick(row, length):
    if not row:
        return ""
    s, l = row.get("Short_Prompt", ""), row.get("Long_Prompt", "")
    return (l or s) if length == "Long" else (s or l)


def _clean(text):
    return text.strip().rstrip(".").strip()


def _dedupe_join(parts):
    seen, out = set(), []
    for p in parts:
        for frag in [f.strip() for f in _clean(p).split(",")]:
            if frag and frag.lower() not in seen:
                seen.add(frag.lower())
                out.append(frag)
    return ", ".join(out)


def _and_list(items):
    items = [i for i in items if i]
    if len(items) <= 1:
        return "".join(items)
    return ", ".join(items[:-1]) + " and " + items[-1]


# --------------------------------------------------------------------------- node layout
SLOTS = [("top_1", "Top"), ("top_2", "Top"), ("bottom_1", "Bottom"), ("bottom_2", "Bottom"),
         ("full_body", "Full Body"), ("outerwear", "Outerwear"), ("footwear", "Footwear")]
ACC_COUNT = 6
FACE_FIELDS = [("face_shape", "Face Shape"), ("jaw", "Jaw"), ("cheekbones", "Cheekbones"),
               ("eyebrows", "Eyebrows"), ("eye_shape", "Eye Shape"), ("eye_color", "Eye Color"),
               ("eyelashes", "Eyelashes"), ("nose", "Nose"), ("lips", "Lips"),
               ("complexion", "Complexion"), ("facial_hair", "Facial Hair"),
               ("distinctive_1", "Distinctive"), ("distinctive_2", "Distinctive")]
GENDERS = {"Female": {}, "Male": {}, "Androgynous": {}}


def _section(default=False):
    return ("BOOLEAN", {"default": default, "label_on": "▼ shown", "label_off": "▶ hidden"})


class CharacterAtlas:
    @classmethod
    def INPUT_TYPES(cls):
        a = LIB.attrs
        req = {
            "prompt": ("STRING", {"multiline": True, "default": ""}),
            "mode": (["Custom", "Known Character", "Preset (Actor)"],),
            "known_character": (_opts("known_character", LIB.known),),
            "preset": (_opts("preset", LIB.presets),),
            "recognition": (["Name + Description", "Name Only", "Description Only"],),
            "prompt_length": (["Short", "Long"],),
            "character_position": (["Character, Prompt", "Prompt, Character"],),
            "seed": ("INT", {"default": 0, "min": 0, "max": 0xFFFFFFFFFFFFFFFF,
                             "control_after_generate": True,
                             "tooltip": "Drives every 🎲 Random choice. Fixed seed = same character."}),

            "section_identity": _section(True),
            "archetype": (_opts("archetype", LIB.archetypes),),
            "gender": (_opts("gender", GENDERS),),
            "age": (_opts("age", a.get("Age", {})),),
            "character_type": (_opts("character_type", a.get("Character Type", {})),),
            "height": (_opts("height", a.get("Height", {})),),
            "body_shape": (_opts("body_shape", a.get("Body Shape", {})),),

            "section_face": _section(False),
        }
        for w, t in FACE_FIELDS:
            req[w] = (_opts(w, a.get(t, {})),)

        req["section_expression"] = _section(False)
        req["expression_1"] = (_opts("expression_1", LIB.expressions),)
        req["expression_2"] = (_opts("expression_2", LIB.expressions),)

        req["section_hair"] = _section(True)
        req["hair_style"] = (_opts("hair_style", a.get("Hair Style", {})),)
        req["hair_color"] = (_opts("hair_color", a.get("Hair Color", {})),)

        req["section_outfit"] = _section(False)
        for w, slot in SLOTS:
            req[w] = (_opts(w, LIB.outfits.get(slot, {})),)
            req[f"{w}_color"] = (_opts(f"{w}_color", LIB.colors),)
            req[f"{w}_fabric"] = (_opts(f"{w}_fabric", LIB.fabrics_for(slot)),)
            req[f"{w}_pattern"] = (_opts(f"{w}_pattern", LIB.patterns),)
            req[f"{w}_finish"] = (_opts(f"{w}_finish", LIB.finishes),)

        req["section_accessories"] = _section(False)
        for i in range(1, ACC_COUNT + 1):
            req[f"accessory_{i}"] = (_opts(f"accessory_{i}", LIB.accessories),)
            req[f"accessory_{i}_color"] = (_opts(f"accessory_{i}_color", LIB.colors),)
            req[f"accessory_{i}_material"] = (_opts(f"accessory_{i}_material", LIB.fabrics_for("Accessory")),)

        req["section_pose"] = _section(False)
        req["pose"] = (_opts("pose", LIB.poses),)
        return {"required": req}

    RETURN_TYPES = ("STRING", "STRING")
    RETURN_NAMES = ("prompt", "character")
    FUNCTION = "build"
    CATEGORY = "Style Atlas/Prompt"

    # ---------------------------------------------------------------- pieces
    def _subject(self, k, L):
        gender, age = k.get("gender", NONE), k.get("age", NONE)
        ctype = k.get("character_type", NONE)
        noun = {"Female": "woman", "Male": "man", "Androgynous": "androgynous person"}.get(gender, "")
        age_word = "" if age == NONE else _pick(LIB.attrs["Age"].get(age), L)
        if age in ("Child", "Teen") and gender in ("Female", "Male"):
            noun = {"Female": "girl", "Male": "boy"}[gender]
            age_word = "teenage" if age == "Teen" else ""
        type_word = "" if ctype in (NONE, "Human") else _pick(LIB.attrs["Character Type"].get(ctype), L)
        if not noun and age_word and not type_word:
            noun = "person"
        return " ".join(x for x in (age_word, type_word, noun) if x)

    def _attr(self, t, key, L):
        return "" if key == NONE else _pick(LIB.attrs.get(t, {}).get(key), L)

    def _garment(self, k, w, slot, L):
        row = LIB.outfits.get(slot, {}).get(k.get(w, NONE))
        if not row:
            return ""
        term = row.get("Garment_Term") or row.get("Short_Prompt") or row["Name"].lower()
        fab_key, pat_key = k.get(f"{w}_fabric", NONE), k.get(f"{w}_pattern", NONE)
        col_key, fin_key = k.get(f"{w}_color", NONE), k.get(f"{w}_finish", NONE)
        fabric = _pick(LIB.fabrics.get(fab_key), L) if fab_key != NONE else row.get("Default_Fabric", "")
        pattern = (LIB.patterns.get(pat_key, {}).get("Prompt", "") if pat_key != NONE
                   else row.get("Default_Pattern", ""))
        color = LIB.colors.get(col_key, {}).get("Prompt", "") if col_key != NONE else ""
        finish = LIB.finishes.get(fin_key, {}).get("Prompt", "") if fin_key != NONE else ""
        phrase = " ".join(x for x in (finish, color, pattern, fabric, term) if x)
        if L == "Long":
            sp, lp = row.get("Short_Prompt", ""), row.get("Long_Prompt", "")
            construction = lp[len(sp):].strip(" ,") if sp and lp.startswith(sp) else ""
            if construction:
                phrase = f"{phrase} {construction}"
        return phrase

    def _accessory(self, k, i, L):
        row = LIB.accessories.get(k.get(f"accessory_{i}", NONE))
        if not row:
            return ""
        col, mat = k.get(f"accessory_{i}_color", NONE), k.get(f"accessory_{i}_material", NONE)
        color = LIB.colors.get(col, {}).get("Prompt", "") if col != NONE else ""
        material = _pick(LIB.fabrics.get(mat), L) if mat != NONE else ""
        return " ".join(x for x in (color, material, row.get("Short_Prompt") or row["Name"].lower()) if x)

    def _base(self, k, L):
        mode, rec = k.get("mode"), k.get("recognition", "Name + Description")
        if mode == "Known Character":
            row = LIB.known.get(k.get("known_character", NONE))
            if not row:
                return ""
            head = f"{row['Character']} from {row['Franchise']}"
        elif mode == "Preset (Actor)":
            row = LIB.presets.get(k.get("preset", NONE))
            if not row:
                return ""
            head = f"{row['Role']} from {row['Franchise']}, played by {row['Actor']}"
        else:
            return None
        text = _clean(_pick(row, L))
        if rec == "Name Only":
            return head
        if rec == "Description Only":
            return text[len(head):].lstrip(", ") if text.startswith(head) else text
        return text

    # ---------------------------------------------------------------- main
    def build(self, prompt, seed=0, **k):
        if LIB.error:
            raise RuntimeError(f"Character Atlas: {LIB.error}")
        k = {name: _resolve(name, v, seed) for name, v in k.items()}
        L = k.get("prompt_length", "Short")
        parts = []

        base = self._base(k, L)
        if base:
            parts.append(base)
        subj = self._subject(k, L)
        if subj:
            parts.append(subj)
        if k.get("archetype", NONE) != NONE:
            parts.append(_pick(LIB.archetypes.get(k["archetype"]), L))

        parts += [self._attr("Height", k.get("height", NONE), L),
                  self._attr("Body Shape", k.get("body_shape", NONE), L)]
        face_detail = False
        for w, t in FACE_FIELDS:
            v = self._attr(t, k.get(w, NONE), L)
            if v.endswith(", clearly defined"):
                v, face_detail = v[: -len(", clearly defined")], True
            parts.append(v)
        if face_detail:
            parts.append("clearly defined facial features")
        for e in ("expression_1", "expression_2"):
            if k.get(e, NONE) != NONE:
                parts.append(_pick(LIB.expressions.get(k[e]), L))
        parts += [self._attr("Hair Color", k.get("hair_color", NONE), L),
                  self._attr("Hair Style", k.get("hair_style", NONE), L)]
        desc = _dedupe_join([p for p in parts if p])

        full = k.get("full_body", NONE) != NONE
        garments = [g for g in (self._garment(k, w, slot, L) for w, slot in SLOTS
                                if not (full and slot in ("Top", "Bottom"))) if g]
        accs = [a for a in (self._accessory(k, i, L) for i in range(1, ACC_COUNT + 1)) if a]
        pose = _clean(_pick(LIB.poses.get(k.get("pose", NONE)), L)) if k.get("pose", NONE) != NONE else ""

        chunks = [desc] if desc else []
        if garments:
            chunks.append("wearing " + _and_list(garments))
        if accs:
            chunks.append("with " + _and_list(accs))
        if pose:
            chunks.append(pose)
        character = ", ".join(c for c in chunks if c)

        user = (prompt or "").strip()
        if k.get("character_position") == "Prompt, Character":
            final = ", ".join(x for x in (user, character) if x)
        else:
            final = ", ".join(x for x in (character, user) if x)
        return (final, character)


NODE_CLASS_MAPPINGS = {"CharacterAtlas": CharacterAtlas}
NODE_DISPLAY_NAME_MAPPINGS = {"CharacterAtlas": "🧑 Character Atlas"}
